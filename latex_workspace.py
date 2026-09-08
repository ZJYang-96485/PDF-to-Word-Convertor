"""Small local Overleaf-style workspace for LaTeX projects.

The workspace keeps the user's source file as the editable authority.  It
compiles a temporary copy so generated files and recovered PDF assets do not
pollute the source folder, displays the resulting PDF beside the source, and
exports the source to either an exact-appearance or editable Word document.
"""

from __future__ import annotations

import base64
import logging
import os
import platform
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import tkinter as tk
import zipfile
from dataclasses import dataclass
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Callable
from converter import (
    VisualFidelityConverter,
    normalise_output_path,
)


LOGGER = logging.getLogger(__name__)
INCLUDE_GRAPHICS_RE = re.compile(
    r"\\includegraphics(?:\s*\[[^]]*\])?\s*\{([^}]+)\}"
)
INCLUDE_GRAPHICS_DETAILS_RE = re.compile(
    r"\\includegraphics\s*(?:\[([^]]*)\])?\s*\{([^}]+)\}"
)


class LatexWorkspaceError(Exception):
    """Expected project or compilation error with a user-facing message."""

    def __init__(self, message: str, user_message: str | None = None) -> None:
        super().__init__(message)
        self.user_message = user_message or message


@dataclass
class CompileResult:
    pdf_path: Path
    log: str
    recovered_assets: list[str]
    build_dir: Path


@dataclass
class ExportResult:
    output_path: str
    mode: str


def _command_path(name: str) -> str | None:
    return shutil.which(name)


def _is_standalone_tex(source: str) -> bool:
    return "\\documentclass" in source and "\\begin{document}" in source


def _graphics_names(source: str) -> list[str]:
    names: list[str] = []
    for raw_name in INCLUDE_GRAPHICS_RE.findall(source):
        name = raw_name.strip()
        if name and name not in names:
            names.append(name)
    return names


def _graphics_widths(source: str) -> dict[str, float]:
    """Return explicit ``includegraphics`` widths in inches when available."""

    widths: dict[str, float] = {}
    for options, raw_name in INCLUDE_GRAPHICS_DETAILS_RE.findall(source):
        match = re.search(r"(?:^|,)\s*width\s*=\s*([0-9.]+)\s*in\b", options)
        if match:
            widths.setdefault(raw_name.strip(), float(match.group(1)))
    return widths


def _find_graphic(source_dir: Path, requested_name: str) -> Path | None:
    requested = Path(requested_name)
    candidates = [source_dir / requested]
    if requested.suffix == "":
        candidates.extend(source_dir / f"{requested_name}{ext}" for ext in (".pdf", ".png", ".jpg", ".jpeg"))
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    return None


def _reference_name_tokens(path: Path) -> set[str]:
    """Normalize a TeX/PDF filename for conservative project matching."""

    return set(re.findall(r"[a-z0-9]+", path.stem.lower()))


def _find_reference_pdf(tex_path: Path) -> Path | None:
    """Find the most likely matching reference PDF beside a source file."""

    exact = tex_path.with_suffix(".pdf")
    if exact.is_file():
        return exact

    source_tokens = _reference_name_tokens(tex_path)
    if not source_tokens:
        return None
    candidates: list[tuple[int, Path]] = []
    for candidate in sorted(tex_path.parent.glob("*.pdf")):
        overlap = len(source_tokens & _reference_name_tokens(candidate))
        if overlap >= 2:
            candidates.append((overlap, candidate))
    if not candidates:
        return None
    candidates.sort(key=lambda item: (-item[0], item[1].name.lower()))
    best_score = candidates[0][0]
    best = [candidate for score, candidate in candidates if score == best_score]
    return best[0] if len(best) == 1 else None


def _copy_existing_graphics(source_dir: Path, build_dir: Path, source: str) -> None:
    for requested_name in _graphics_names(source):
        existing = _find_graphic(source_dir, requested_name)
        if existing is None:
            continue
        destination = build_dir / Path(requested_name)
        destination.parent.mkdir(parents=True, exist_ok=True)
        if not destination.exists():
            shutil.copy2(existing, destination)


def _recover_graphics_from_pdf(
    source: str,
    reference_pdf: Path | None,
    build_dir: Path,
) -> list[str]:
    """Recover missing includegraphics files from a matching reference PDF."""

    missing_names = [
        name
        for name in _graphics_names(source)
        if not _find_graphic(build_dir, name)
    ]
    if not missing_names or reference_pdf is None or not reference_pdf.is_file():
        return []

    try:
        import pymupdf as fitz  # type: ignore[import-not-found]
    except ImportError:
        try:
            import fitz  # type: ignore[import-not-found]
        except ImportError as exc:
            raise LatexWorkspaceError(
                "PyMuPDF is required to recover missing figures from a reference PDF.",
                "Install the project requirements, or place the original figure files beside the .tex file.",
            ) from exc

    recovered: list[str] = []
    try:
        document = fitz.open(str(reference_pdf))
        with document:
            image_entries: list[tuple[int, object, object | None]] = []
            for page in document:
                for image in page.get_images(full=True):
                    xref = int(image[0])
                    if any(entry[0] == xref for entry in image_entries):
                        continue
                    rectangles = page.get_image_rects(xref)
                    image_entries.append((xref, page, rectangles[0] if rectangles else None))

            if len(image_entries) < len(missing_names):
                raise LatexWorkspaceError(
                    f"The reference PDF contains {len(image_entries)} embedded figure(s), but the source needs {len(missing_names)}.",
                    "The source refers to figures that could not all be recovered from the matching PDF.",
                )

            for requested_name, (xref, page, rectangle) in zip(missing_names, image_entries):
                destination = build_dir / Path(requested_name)
                destination.parent.mkdir(parents=True, exist_ok=True)
                pixmap = fitz.Pixmap(document, xref)
                try:
                    samples = pixmap.samples
                    if samples and max(samples) != min(samples):
                        # The source usually specifies .jpg.  Pixmap.save()
                        # writes the format implied by that extension.
                        pixmap.save(str(destination))
                    elif rectangle is not None:
                        # Some LaTeX/PDF pipelines leave a transparent image
                        # XObject in the PDF while the visible figure is
                        # rendered into the page.  Crop that page rectangle
                        # instead of exporting a solid black placeholder.
                        rendered = page.get_pixmap(
                            matrix=fitz.Matrix(4, 4),
                            clip=rectangle,
                            alpha=False,
                        )
                        try:
                            rendered.save(str(destination))
                        finally:
                            rendered = None
                    else:
                        raise LatexWorkspaceError(
                            f"Figure {requested_name} has no visible PDF image data.",
                            "The matching PDF contains a figure that could not be recovered.",
                        )
                finally:
                    pixmap = None
                recovered.append(destination.name)
    except LatexWorkspaceError:
        raise
    except Exception as exc:
        LOGGER.exception("Could not recover figures from %s", reference_pdf)
        raise LatexWorkspaceError(
            f"Could not recover figures from {reference_pdf}: {exc}",
            "The matching PDF was found, but its figures could not be extracted for compilation.",
        ) from exc
    return recovered


def _wrapper_source(body_name: str) -> str:
    return f"""\\documentclass[10pt]{{exam}}
\\usepackage[margin=1in]{{geometry}}
\\usepackage{{amsmath,amssymb,enumitem,graphicx}}
\\pointformat{{[\\thepoints]}}
\\printanswers
\\begin{{document}}
\\input{{{body_name}}}
\\end{{document}}
"""


_TEX_INPUT_RE = re.compile(r"\\(?:input|include)\s*\{([^}]+)\}")


def _active_tex_references(source: str) -> list[str]:
    """Return local file references from non-commented TeX lines."""

    active_source = "\n".join(
        line for line in source.splitlines() if not line.lstrip().startswith("%")
    )
    return _TEX_INPUT_RE.findall(active_source)


def _find_master_template(tex_path: Path, source: str) -> Path | None:
    """Find a local master file that includes the selected source fragment.

    A fragment is often compiled by VSCode through a project master such as
    ``main.tex``.  Reusing that master is important because its packages,
    fonts, and custom style files are part of the document's visual contract.
    Only an exact include of the selected source is accepted, so an unrelated
    TeX file in the same folder cannot silently change the result.
    """

    if _is_standalone_tex(source):
        return None
    target_names = {tex_path.name, tex_path.stem}
    for candidate in sorted(tex_path.parent.glob("*.tex")):
        if candidate.resolve() == tex_path.resolve():
            continue
        try:
            candidate_source = candidate.read_text(encoding="utf-8")
        except OSError:
            continue
        if not _is_standalone_tex(candidate_source):
            continue
        for reference in _active_tex_references(candidate_source):
            reference_name = Path(reference.strip()).name
            if reference_name in target_names or Path(reference_name).stem in target_names:
                return candidate
    return None


def _resolve_tex_reference(base_dir: Path, reference: str) -> Path | None:
    """Resolve a local ``input``/``include`` reference, with TeX extensions."""

    requested = Path(reference.strip())
    if requested.is_absolute():
        return None
    candidates = [base_dir / requested]
    if requested.suffix == "":
        candidates.extend((base_dir / f"{requested}{extension}") for extension in (".tex", ".sty"))
    return next((candidate for candidate in candidates if candidate.is_file()), None)


def _template_wrapper_source(
    template_path: Path,
    body_name: str,
    project_dir: Path,
) -> tuple[str, list[str]]:
    """Build a wrapper from a project's master preamble and copy its styles."""

    template = template_path.read_text(encoding="utf-8")
    preamble, _separator, _document_body = template.partition(r"\begin{document}")
    warnings: list[str] = [f"Using project style template: {template_path.name}"]
    pending = _active_tex_references(preamble)
    visited: set[Path] = set()
    missing: list[str] = []
    while pending:
        reference = pending.pop(0)
        resolved = _resolve_tex_reference(template_path.parent, reference)
        if resolved is None:
            missing.append(reference)
            continue
        resolved = resolved.resolve()
        if resolved in visited:
            continue
        visited.add(resolved)
        relative = resolved.relative_to(template_path.parent.resolve())
        destination = project_dir / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(resolved, destination)
        try:
            support_source = resolved.read_text(encoding="utf-8")
        except OSError:
            support_source = ""
        pending.extend(_active_tex_references(support_source))

    for reference in missing:
        escaped = re.escape(reference)
        preamble, replacements = re.subn(
            rf"^[ \t]*\\(?:input|include)\s*\{{{escaped}\}}[ \t]*(?:%[^\n]*)?(?:\n|$)",
            f"% PDF-to-Word: omitted missing local style file {reference}\n",
            preamble,
            flags=re.MULTILINE,
        )
        if replacements:
            warnings.append(f"Omitted missing local style file: {reference}")

    wrapper = f"{preamble}\n\\begin{{document}}\n\\input{{{body_name}}}\n\\end{{document}}\n"
    return wrapper, warnings


_TEX4HT_DISPLAY_ALIGNMENT_RE = re.compile(
    r"\$\$\s*\\begin\{aligned\}(.*?)\\end\{aligned\}\s*\$\$",
    re.DOTALL,
)
_TEX4HT_STANDALONE_ALIGNMENT_RE = re.compile(
    r"\\begin\{(align\*|eqnarray\*|gather\*)\}(.*?)\\end\{\1\}",
    re.DOTALL,
)
_TEX4HT_INLINE_DOLLAR_RE = re.compile(r"(?<!\$)\$(?!\$)([^$\n]+?)\$(?!\$)")
_TEX_SIZE_DELIMITER_RE = re.compile(
    r"\\(?:big|Big|bigg|Bigg|bigl|Bigl|biggl|Biggl|bigr|Bigr|biggr|Biggr)\s*(\\[{}]|[()[\]])"
)
_TEX_ROMAN_OPERATOR_RE = re.compile(r"\\(exp|ln)(?![A-Za-z])")
_TEX_ENV_TOKEN_RE = re.compile(r"\\(begin|end)\s*\{(parts|solution)\}")
_BOX_MARKER_RE = re.compile(r"PDFBox[SE]\d{3}")
_EQUALITY_MARKER_RE = re.compile(r"PDFEq\d{3}")
_ROMAN_MARKER_RE = re.compile(r"PDFRoman[SE]\d{3}")


def _mark_boxed_math(source: str) -> str:
    """Add private markers so boxed formulas can be restored as Word boxes.

    TeX4ht and LibreOffice preserve the formula itself as editable OMML, but
    they otherwise drop ``\\boxed{...}``.  A temporary text marker travels
    through the math conversion and is removed after the DOCX is created.
    """

    output: list[str] = []
    cursor = 0
    marker_number = 1
    while True:
        start_token = source.find(r"\boxed{", cursor)
        if start_token < 0:
            output.append(source[cursor:])
            break
        output.append(source[cursor:start_token])
        content_start = start_token + len(r"\boxed{")
        depth = 1
        index = content_start
        while index < len(source) and depth:
            character = source[index]
            escaped = index > 0 and source[index - 1] == "\\"
            if character == "{" and not escaped:
                depth += 1
            elif character == "}" and not escaped:
                depth -= 1
            index += 1
        if depth:
            output.append(source[start_token:])
            break
        content = source[content_start : index - 1]
        start_marker = f"PDFBoxS{marker_number:03d}"
        end_marker = f"PDFBoxE{marker_number:03d}"
        output.append(
            r"\boxed{"
            + rf"\text{{{start_marker}}}\,"
            + content
            + rf"\,\text{{{end_marker}}}"
            + "}"
        )
        marker_number += 1
        cursor = index
    return "".join(output)


def _flatten_nested_exam_solutions(source: str) -> str:
    """Keep solutions nested in ``exam`` parts visible in TeX4ht output.

    TeX4ht's ODT backend can drop an ``exam`` solution when it occurs inside
    the outer ``parts`` list.  For the Word-export copy only, flatten that
    solution into ordinary editable text, equations, and lists while keeping
    it in source order; this avoids the malformed nested ODT list.  The
    user's source and normal PDF compilation are not changed.
    """

    solution_block = re.compile(
        r"\\begin\s*\{solution\}(.*?)\\end\s*\{solution\}",
        re.DOTALL,
    )
    output: list[str] = []
    cursor = 0
    parts_depth = 0

    def update_parts_depth(fragment: str) -> None:
        nonlocal parts_depth
        for token in _TEX_ENV_TOKEN_RE.finditer(fragment):
            kind, environment = token.groups()
            if environment != "parts":
                continue
            parts_depth += 1 if kind == "begin" else -1
            parts_depth = max(0, parts_depth)

    for match in solution_block.finditer(source):
        prefix = source[cursor : match.start()]
        output.append(prefix)
        update_parts_depth(prefix)

        if parts_depth == 0:
            output.append(match.group(0))
        else:
            body = match.group(1)
            body = body.replace(
                r"\begin{parts}", r"\begin{enumerate}[label=(\alph*)]"
            )
            body = body.replace(r"\end{parts}", r"\end{enumerate}")
            body = re.sub(
                r"\\part(?:\s*\[[^]]*\])?",
                lambda _match: r"\item",
                body,
            )
            output.append(r"\par\noindent\textbf{Solution:}\par" + body + r"\par")
            update_parts_depth(match.group(0))
        cursor = match.end()
    output.append(source[cursor:])
    return "".join(output)


def _normalise_size_delimiters(source: str) -> str:
    r"""Export ``\big``-style delimiters as native scalable delimiters.

    TeX4ht's ODT backend can emit an empty MathML delimiter for commands such
    as ``\big(``, leaving the fraction or expression outside the delimiter in
    the resulting Word equation.  ``\left``/``\right`` uses the same visible
    parentheses while producing one editable OMML delimiter around the
    expression.
    """

    def replace(match: re.Match[str]) -> str:
        delimiter = match.group(1)
        closing = delimiter in (")", "]", r"\}")
        return (r"\right" if closing else r"\left") + delimiter

    return _TEX_SIZE_DELIMITER_RE.sub(replace, source)


def _mark_roman_math(source: str) -> str:
    r"""Mark source-level roman math spans so Word can restore their style.

    The ODT route preserves the equation structure but loses the distinction
    between ordinary math italic and ``\mathrm``.  Text markers survive that
    conversion, allowing the DOCX postprocessor to apply native OMML roman
    styling only to the requested spans.  The parser is deliberately small:
    it only needs to find balanced ``\mathrm{...}`` groups and the common
    operator commands ``\exp`` and ``\ln``.
    """

    output: list[str] = []
    cursor = 0
    marker_number = 1

    def marker_pair() -> tuple[str, str]:
        nonlocal marker_number
        start = f"PDFRomanS{marker_number:03d}"
        end = f"PDFRomanE{marker_number:03d}"
        marker_number += 1
        return start, end

    while cursor < len(source):
        if source.startswith(r"\mathrm", cursor):
            brace = cursor + len(r"\mathrm")
            while brace < len(source) and source[brace].isspace():
                brace += 1
            if brace < len(source) and source[brace] == "{":
                depth = 1
                index = brace + 1
                while index < len(source) and depth:
                    character = source[index]
                    escaped = index > 0 and source[index - 1] == "\\"
                    if character == "{" and not escaped:
                        depth += 1
                    elif character == "}" and not escaped:
                        depth -= 1
                    index += 1
                if depth == 0:
                    start_marker, end_marker = marker_pair()
                    content = source[brace + 1 : index - 1]
                    output.append(source[cursor : brace + 1])
                    output.append(rf"\text{{{start_marker}}}")
                    output.append(content)
                    output.append(rf"\text{{{end_marker}}}")
                    output.append("}")
                    cursor = index
                    continue

        operator = _TEX_ROMAN_OPERATOR_RE.match(source, cursor)
        if operator is not None:
            start_marker, end_marker = marker_pair()
            output.append(
                rf"\mathrm{{\text{{{start_marker}}}{operator.group(1)}\text{{{end_marker}}}}}"
            )
            cursor = operator.end()
            continue

        output.append(source[cursor])
        cursor += 1

    return "".join(output)


def _tex4ht_source(source: str) -> str:
    """Make alignment-heavy math safe for TeX4ht's ODT backend.

    TeX4ht correctly converts ordinary LaTeX math to MathML, which LibreOffice
    then converts to native Word OMML.  Its ODT backend does not handle the
    alignment tabs in ``aligned``/``align``/``eqnarray`` reliably, though: the
    tabs can become stray glyphs or push a formula outside the page width.
    Keep every equation and line break, but export each aligned row as a
    separate display equation.  This avoids TeX4ht's malformed ODT tables and
    keeps long rows inside the Word page width.  The user's source and the PDF
    preview remain unchanged.
    """

    equality_marker_number = 1

    def rows(body: str) -> list[str]:
        nonlocal equality_marker_number
        cleaned_rows: list[str] = []
        for row in re.split(r"\\\\(?:\s*\[[^]]*\])?", body):
            cleaned = row.replace("&", "").strip()
            if not cleaned:
                continue
            # TeX4ht treats an equals sign at the very start of a standalone
            # equation as alignment metadata and can drop it during the ODT
            # conversion.  Put a temporary text marker before it so the
            # equality is no longer the first token; the marker is removed
            # after LibreOffice creates the editable Word math.
            if cleaned.startswith("="):
                marker = f"PDFEq{equality_marker_number:03d}"
                equality_marker_number += 1
                cleaned = rf"\text{{{marker}}}\mathrel{{=}}" + cleaned[1:]
            cleaned_rows.append(cleaned)
        return cleaned_rows

    def replace_display_alignment(match: re.Match[str]) -> str:
        return "\n".join(
            f"\\begin{{equation*}}\n{row}\n\\end{{equation*}}" for row in rows(match.group(1))
        )

    def replace_standalone_alignment(match: re.Match[str]) -> str:
        return "\n".join(
            f"\\begin{{equation*}}\n{row}\n\\end{{equation*}}" for row in rows(match.group(2))
        )

    source = _flatten_nested_exam_solutions(source)
    source = _normalise_size_delimiters(source)
    source = _mark_boxed_math(source)
    # TeX4ht's MathML writer can emit malformed fragments for thin-space
    # commands inside long ``\mathrm`` subscripts.  Ordinary TeX spaces keep
    # the same visible wording and produce a valid editable Word formula.
    source = source.replace(r"\mathrm{out\,of\,engine}", r"{out\ of\ engine}")
    source = source.replace(r"\mathrm{in\,to\,engine}", r"{in\ to\ engine}")
    source = _mark_roman_math(source)
    source = _TEX4HT_DISPLAY_ALIGNMENT_RE.sub(replace_display_alignment, source)
    source = _TEX4HT_STANDALONE_ALIGNMENT_RE.sub(replace_standalone_alignment, source)
    if r"\begin{enumerate}" in source and r"\labelenumi" not in source:
        source = source.replace(
            r"\begin{enumerate}",
            r"\renewcommand{\labelenumi}{\arabic{enumi}.}\begin{enumerate}",
        )

    def preserve_inline_word_space(match: re.Match[str]) -> str:
        following = source[match.end() :]
        next_significant = re.match(r"\s*(\S)", following)
        if next_significant and (next_significant.group(1).isalnum() or next_significant.group(1) == "\\"):
            return match.group(0) + r"\ "
        return match.group(0)

    return _TEX4HT_INLINE_DOLLAR_RE.sub(preserve_inline_word_space, source)


def _tool_path(names: tuple[str, ...], extra_paths: tuple[str, ...] = ()) -> str | None:
    for name in names:
        found = _command_path(name)
        if found:
            return found
    for candidate in extra_paths:
        if Path(candidate).is_file():
            return candidate
    return None


def _filter_resolved_reference_warnings(log: str, aux_path: Path) -> str:
    """Hide first-pass cross-reference warnings resolved by the final pass."""

    if not aux_path.is_file():
        return log
    try:
        aux = aux_path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return log
    resolved = set(re.findall(r"\\newlabel\{([^}]+)\}", aux))
    if not resolved:
        return log

    reference_warning = re.compile(r"Reference [`']([^`']+)[`'] .*undefined")
    filtered: list[str] = []
    removed_reference_warning = False
    for line in log.splitlines():
        match = reference_warning.search(line)
        if match and match.group(1) in resolved:
            removed_reference_warning = True
            continue
        filtered.append(line)
    if removed_reference_warning and not any(reference_warning.search(line) for line in filtered):
        filtered = [line for line in filtered if "There were undefined references." not in line]
    return "\n".join(filtered)


def _repair_odt_content_xml(data: bytes, etree: object) -> bytes:
    """Repair malformed list boundaries emitted by TeX4ht's ODT backend."""

    try:
        etree.fromstring(data)
        return data
    except etree.XMLSyntaxError:
        pass

    # TeX4ht can emit a pair of stray paragraph tags around a gather/aligned
    # display.  Removing them exposes the intended list boundaries.
    repaired = re.sub(
        rb"\s*</text:p>\s*<text:p\b[^>]*></mtable>\s*</text:p>",
        b"",
        data,
    )

    # A gathered display can leave raw MathML closing tags inside the next
    # display paragraph.  Close that paragraph before the following one so
    # the later question/list content remains inside office:text.
    stray_math = re.compile(
        rb"(</draw:frame>)(?:(?!<text:p\b).)*?(?=<text:p\b)",
        re.DOTALL,
    )

    def close_stray_math(match: re.Match[bytes]) -> bytes:
        block = match.group(0)
        if any(marker in block for marker in (b"</mrow", b"</msub", b"<mi")):
            return match.group(1) + b"</text:p>\n"
        return block

    repaired = stray_math.sub(close_stray_math, repaired)

    # When a solution is nested in an exam parts list, TeX4ht leaves the
    # following question as raw text between list tags.  Put that text back
    # into a paragraph and close the stale outer list before it.
    orphan_text = re.compile(
        rb"(</text:list>)\s+(?=[^<\s])(.*?)(?=<text:list\b)",
        re.DOTALL,
    )
    repaired, orphan_count = orphan_text.subn(
        lambda match: (
            match.group(1)
            + b"\n</text:list-item></text:list>\n"
            + b'<text:p text:style-name="Text-body">'
            + match.group(2)
            + b"</text:p>\n"
        ),
        repaired,
    )

    try:
        etree.fromstring(repaired)
        return repaired
    except etree.XMLSyntaxError:
        if orphan_count:
            closing = b"</text:list-item></text:list>\n"
            office_text = b"</office:text>"
            position = repaired.rfind(office_text)
            if position >= 0:
                closed = repaired[:position] + closing + repaired[position:]
                try:
                    etree.fromstring(closed)
                    return closed
                except etree.XMLSyntaxError:
                    repaired = closed

    parser = etree.XMLParser(recover=True, huge_tree=True)
    root = etree.fromstring(repaired, parser)
    if root is None:
        raise LatexWorkspaceError(
            "TeX4ht produced an unrecoverable content.xml file.",
            "The editable Word export could not repair the intermediate document structure.",
        )
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _repair_odt_manifest(odt_path: Path) -> Path:
    """Repair TeX4ht ODT metadata/XML before LibreOffice imports it.

    Older TeX4ht releases can emit a malformed ``content.xml`` for nested
    ``exam`` parts/lists.  LibreOffice refuses that package even though the
    PDF compilation succeeded.  Recovering only that XML file keeps the text
    and editable math while allowing LibreOffice to open the document.
    """

    repaired = odt_path.with_name(odt_path.stem + "-fixed.odt")
    try:
        from lxml import etree  # type: ignore[import-not-found]
    except ImportError as exc:
        raise LatexWorkspaceError(
            "lxml is required to repair TeX4ht's ODT output.",
            "Install the project requirements, then try editable Word export again.",
        ) from exc

    with zipfile.ZipFile(odt_path, "r") as source, zipfile.ZipFile(
        repaired, "w", compression=zipfile.ZIP_DEFLATED
    ) as destination:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "META-INF/manifest.xml":
                manifest = data.decode("utf-8")
                manifest = manifest.replace(r"\image:mimejpg ", "image/jpeg")
                manifest = manifest.replace(r"\image:mimepng ", "image/png")
                manifest = manifest.replace(r"\image:mimegif ", "image/gif")
                manifest = manifest.replace(r"\image:mimebmp ", "image/bmp")
                data = manifest.encode("utf-8")
            elif item.filename == "content.xml":
                data = _repair_odt_content_xml(data, etree)
            destination.writestr(item, data)
    return repaired


class LatexSourceDocxConverter:
    """Convert compiled LaTeX source to editable Word math and text.

    The source is first passed through TeX4ht's ODT backend.  LibreOffice then
    converts the ODT into DOCX, preserving the equations as Word OMML rather
    than flattening them into page screenshots.
    """

    def __init__(self, compile_result: CompileResult, tex_path: Path, source_text: str) -> None:
        self.compile_result = compile_result
        self.tex_path = tex_path
        self.source_text = source_text

    def convert_pdf(
        self,
        pdf_path: Path,
        output_path: Path,
        *,
        overwrite: bool = False,
        status_callback: Callable[[str], None] | None = None,
    ) -> str:
        output = normalise_output_path(pdf_path, output_path)
        if output.exists() and not overwrite:
            raise LatexWorkspaceError(
                f"The output already exists: {output}",
                "Choose a different Word filename or approve replacing the existing file.",
            )
        if not output.parent.is_dir() or not os.access(output.parent, os.W_OK):
            raise LatexWorkspaceError(
                f"The output folder is not writable: {output.parent}",
                "The Word file could not be saved. Check the selected output folder.",
            )

        make4ht = _tool_path(
            ("make4ht",),
            ("/Library/TeX/texbin/make4ht", "/usr/local/bin/make4ht"),
        )
        soffice = _tool_path(
            ("soffice", "libreoffice"),
            (
                "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                r"C:\\Program Files\\LibreOffice\\program\\soffice.exe",
                r"C:\\Program Files (x86)\\LibreOffice\\program\\soffice.exe",
            ),
        )
        if make4ht is None:
            raise LatexWorkspaceError(
                "TeX4ht make4ht was not found.",
                "Editable LaTeX-to-Word export requires TeX Live or MiKTeX with TeX4ht (make4ht).",
            )
        if soffice is None:
            raise LatexWorkspaceError(
                "LibreOffice soffice was not found.",
                "Editable LaTeX-to-Word export requires LibreOffice. Install it, restart the app, and try again.",
            )

        project_dir = self.compile_result.build_dir.parent / "project"
        main_name = self.compile_result.pdf_path.with_suffix(".tex").name
        compiled_main = project_dir / main_name
        if not compiled_main.is_file():
            raise LatexWorkspaceError(
                f"The compiled LaTeX project is incomplete: {compiled_main}",
                "Compile the LaTeX source again before exporting editable Word.",
            )

        export_root = Path(tempfile.mkdtemp(prefix="pdf_to_word_tex4ht_"))
        tex4ht_project = export_root / "project"
        odt_dir = export_root / "odt"
        docx_dir = export_root / "docx"
        odt_dir.mkdir()
        docx_dir.mkdir()
        shutil.copytree(project_dir, tex4ht_project)

        if _is_standalone_tex(self.source_text):
            (tex4ht_project / main_name).write_text(
                _tex4ht_source(self.source_text), encoding="utf-8"
            )
        else:
            body_path = tex4ht_project / "source_body.tex"
            body_path.write_text(_tex4ht_source(self.source_text), encoding="utf-8")

        self._notify(status_callback, "Converting LaTeX source to editable document structure...")
        make_result = subprocess.run(
            [make4ht, "-a", "warning", "-f", "odt", "-d", str(odt_dir), str(tex4ht_project / main_name)],
            cwd=str(tex4ht_project),
            capture_output=True,
            text=True,
            check=False,
        )
        odt_path = odt_dir / Path(main_name).with_suffix(".odt").name
        make_log = "\n".join(part for part in (make_result.stdout, make_result.stderr) if part)
        if not odt_path.is_file():
            raise LatexWorkspaceError(
                f"TeX4ht could not create an ODT file:\n{make_log}",
                "TeX4ht could not convert this LaTeX source. Review the compiler installation and source log.",
            )

        odt_for_import = _repair_odt_manifest(odt_path)
        self._notify(status_callback, "Converting editable equations to Word format...")
        office_result = subprocess.run(
            [soffice, "--headless", "--convert-to", "docx", "--outdir", str(docx_dir), str(odt_for_import)],
            cwd=str(docx_dir),
            capture_output=True,
            text=True,
            check=False,
        )
        generated_docx = docx_dir / odt_for_import.with_suffix(".docx").name
        office_log = "\n".join(part for part in (office_result.stdout, office_result.stderr) if part)
        if office_result.returncode != 0 or not generated_docx.is_file():
            raise LatexWorkspaceError(
                f"LibreOffice could not create DOCX:\n{office_log}",
                "LibreOffice could not finish the editable Word export. Make sure LibreOffice is installed and try again.",
            )

        try:
            shutil.copy2(generated_docx, output)
        except OSError as exc:
            raise LatexWorkspaceError(
                f"Could not save the DOCX to {output}: {exc}",
                "The editable Word file could not be saved to the selected location.",
            ) from exc
        self._postprocess_docx(output, project_dir)
        return str(output)

    @staticmethod
    def _notify(callback: Callable[[str], None] | None, message: str) -> None:
        if callback is not None:
            try:
                callback(message)
            except Exception:
                LOGGER.exception("Status callback failed.")

    def _postprocess_docx(self, docx_path: Path, project_dir: Path) -> None:
        """Restore source-level structure that ODT cannot represent directly."""

        try:
            from docx import Document  # type: ignore[import-not-found]
        except ImportError as exc:
            raise LatexWorkspaceError(
                "python-docx is required to finish the editable DOCX structure.",
                "Install the project requirements, then try editable Word export again.",
            ) from exc

        document = Document(str(docx_path))
        # TeX4ht's ``dt`` style is bold by default.  It is useful for the
        # source list indentation, but it also becomes the inherited style for
        # OMML equation runs, which do not receive the direct ``w:b=0`` that
        # ordinary text runs receive below.  Make the style itself regular so
        # unbolded equations stay regular; explicit LaTeX \textbf remains a
        # direct run-level override.
        document.styles["dt"].font.bold = False  # type: ignore[attr-defined]
        self._clean_header(document)
        self._restore_question_points(document)
        self._clean_leaked_part_points(document)
        self._restore_part_points(document)
        self._restore_leading_equals(document)
        self._restore_math_roman_styles(document)
        self._restore_boxed_equations(document)
        self._restore_list_labels(document)
        self._restore_solution_frames(document)
        self._add_page_numbers(document)
        self._insert_figures(document, project_dir)
        document.save(str(docx_path))

    @staticmethod
    def _set_paragraph_border(paragraph: object, sides: tuple[str, ...]) -> None:
        from docx.oxml import OxmlElement  # type: ignore[import-not-found]
        from docx.oxml.ns import qn  # type: ignore[import-not-found]

        paragraph_element = paragraph._p  # type: ignore[attr-defined]
        properties = paragraph_element.get_or_add_pPr()
        border = properties.find(qn("w:pBdr"))
        if border is None:
            border = OxmlElement("w:pBdr")
            properties.append(border)
        for side in sides:
            element = border.find(qn(f"w:{side}"))
            if element is None:
                element = OxmlElement(f"w:{side}")
                border.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "6")
            element.set(qn("w:space"), "6")
            element.set(qn("w:color"), "000000")

    @staticmethod
    def _delete_paragraph(paragraph: object) -> None:
        element = paragraph._element  # type: ignore[attr-defined]
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    @staticmethod
    def _paragraph_has_embedded_object(paragraph: object) -> bool:
        element = paragraph._p  # type: ignore[attr-defined]
        return (
            bool(element.xpath(".//w:drawing"))
            or bool(element.xpath(".//w:pict"))
            or bool(element.xpath(".//m:oMath"))
            or bool(element.xpath(".//m:oMathPara"))
        )

    def _clean_header(self, document: object) -> None:
        """Turn TeX4ht's underscore rules into real Word borders."""

        paragraphs = list(document.paragraphs)  # type: ignore[attr-defined]
        first_question = next(
            (
                index
                for index, paragraph in enumerate(paragraphs)
                if re.match(r"^\d+\.\s*(?:\[\d+\s+points?\])?(?:\s|$)", paragraph.text.strip())
            ),
            len(paragraphs),
        )
        for index, paragraph in enumerate(paragraphs[:first_question]):
            if "_" in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace("_", "")
                self._set_paragraph_border(paragraph, ("bottom",))
            elif not paragraph.text.strip() and not self._paragraph_has_embedded_object(paragraph):
                self._delete_paragraph(paragraph)

        tables = list(document.tables)  # type: ignore[attr-defined]
        if tables:
            table = tables[0]
            header_source = re.split(r"\\begin\s*\{questions\}", self.source_text, maxsplit=1)[0]
            header_has_vertical_rules = bool(
                re.search(r"\\begin\s*\{tabular\}\s*\{[^}]*\|", header_source)
            )
            for row in list(table.rows):
                if not any(cell.text.strip() for cell in row.cells):
                    row._tr.getparent().remove(row._tr)
            from docx.oxml import OxmlElement  # type: ignore[import-not-found]
            from docx.oxml.ns import qn  # type: ignore[import-not-found]

            for row in table.rows:
                for cell_index, cell in enumerate(row.cells[:-1]):
                    properties = cell._tc.get_or_add_tcPr()
                    borders = properties.find(qn("w:tcBorders"))
                    if header_has_vertical_rules and borders is None:
                        borders = OxmlElement("w:tcBorders")
                        properties.append(borders)
                    if borders is None:
                        continue
                    right = borders.find(qn("w:right"))
                    if header_has_vertical_rules and right is None:
                        right = OxmlElement("w:right")
                        borders.append(right)
                    if header_has_vertical_rules and right is not None:
                        right.set(qn("w:val"), "single")
                        right.set(qn("w:sz"), "4")
                        right.set(qn("w:space"), "0")
                        right.set(qn("w:color"), "000000")
                    elif right is not None:
                        borders.remove(right)

    @staticmethod
    def _normalise_question_text(text: str) -> str:
        text = re.sub(r"(?m)%.*$", " ", text)
        text = re.sub(r"\\(?:ref|autoref|pageref)\s*\{[^}]*\}", " ", text)
        text = re.sub(r"\$\$.*?\$\$|\$.*?\$", " ", text, flags=re.DOTALL)
        text = re.sub(r"\\[A-Za-z@]+\*?(?:\s*\[[^]]*\])?", " ", text)
        text = re.sub(r"[{}]", " ", text)
        return " ".join(text.split())

    def _question_specs(self) -> list[tuple[str | None, str]]:
        matches = list(re.finditer(r"\\question\s*(?:\[(\d+)\])?", self.source_text))
        specs: list[tuple[str | None, str]] = []
        for index, match in enumerate(matches):
            end = matches[index + 1].start() if index + 1 < len(matches) else len(self.source_text)
            body = self.source_text[match.end() : end]
            lead = self._normalise_question_text(body)
            specs.append((match.group(1), " ".join(lead.split()[:8])))
        return specs

    def _restore_question_points(self, document: object) -> None:
        """Restore question labels inline with their editable question text."""

        paragraphs = list(document.paragraphs)  # type: ignore[attr-defined]
        specs = self._question_specs()
        for question_number, (points, lead) in enumerate(specs, start=1):
            if not lead:
                continue

            def question_body(text: str) -> str:
                normalised = self._normalise_question_text(text)
                return re.sub(
                    r"^\d+\.\s*(?:\[\d+\s+points?\])?\s*",
                    "",
                    normalised,
                )

            candidate_index = next(
                (
                    index
                    for index, paragraph in enumerate(paragraphs)
                    if question_body(paragraph.text).startswith(lead)
                ),
                None,
            )
            if candidate_index is None:
                continue

            label = f"{question_number}."
            if points:
                label += f" [{points} points]"

            marker = paragraphs[candidate_index - 1] if candidate_index > 0 else None
            if (
                marker is not None
                and marker.style
                and marker.style.name == "Inside-enumerate"
                and not marker.text.strip()
            ):
                self._delete_paragraph(marker)
                paragraphs.remove(marker)
                candidate_index -= 1

            paragraph = paragraphs[candidate_index]
            previous_index = candidate_index - 1
            while previous_index >= 0 and not paragraphs[previous_index].text.strip():
                previous_index -= 1
            if (
                previous_index >= 0
                and re.fullmatch(r"\d+\.\s*(?:\[\d+\s+points?\])?", paragraphs[previous_index].text.strip())
            ):
                old_label = paragraphs[previous_index]
                self._delete_paragraph(old_label)
                paragraphs.pop(previous_index)
                candidate_index -= 1
                paragraph = paragraphs[candidate_index]

            label_pattern = re.compile(r"^\d+\.\s*(?:\[\d+\s+points?\])?\s*")
            first_text_run = next((run for run in paragraph.runs if run.text), None)
            if first_text_run is not None and label_pattern.match(first_text_run.text or ""):
                first_text_run.text = label_pattern.sub(label + " ", first_text_run.text, count=1)
            else:
                prefix_run = paragraph.add_run(label + " ")
                prefix_run.bold = False
                prefix_element = prefix_run._r
                prefix_element.getparent().remove(prefix_element)
                insert_at = 1 if paragraph._p.pPr is not None else 0
                paragraph._p.insert(insert_at, prefix_element)

            paragraph.style = document.styles["dt"]  # type: ignore[attr-defined]
            for run in paragraph.runs:
                if run.bold is None:
                    run.bold = False

    @staticmethod
    def _clean_leaked_part_points(document: object) -> None:
        """Remove question totals that TeX4ht incorrectly copies onto ``(a)``."""

        leaked = re.compile(r"^\[\d+ points\]\(([a-z])\)$")
        for paragraph in document.paragraphs:  # type: ignore[attr-defined]
            match = leaked.fullmatch(paragraph.text.strip())
            if match:
                paragraph.text = f"({match.group(1)})"

    def _source_question_part_points(self) -> list[str | None]:
        token_re = re.compile(
            r"\\begin\s*\{solution\}|\\end\s*\{solution\}"
            r"|\\part\s*(?:\[(\d+)\])?"
        )
        solution_depth = 0
        points: list[str | None] = []
        for match in token_re.finditer(self.source_text):
            token = match.group(0)
            if token.startswith(r"\begin"):
                solution_depth += 1
            elif token.startswith(r"\end"):
                solution_depth = max(0, solution_depth - 1)
            elif solution_depth == 0:
                points.append(match.group(1))
        return points

    def _restore_part_points(self, document: object) -> None:
        """Restore point totals on top-level exam parts and keep labels inline."""

        points = self._source_question_part_points()
        paragraphs = [
            paragraph
            for paragraph in document.paragraphs  # type: ignore[attr-defined]
            if re.fullmatch(
                r"(?:\[\d+\s+points?\]\s*)?\([a-z]\)(?:\s*\[\d+\s+points?\])?",
                paragraph.text.strip(),
            )
        ]
        for paragraph, value in zip(paragraphs, points):
            if not value:
                continue
            label = re.search(r"\([a-z]\)", paragraph.text.strip())
            if label is None:
                continue
            rendered_label = f"{label.group(0)} [{value} points]"
            paragraph.text = rendered_label
            for run in paragraph.runs:
                run.bold = False

            # A minipage after \part is emitted by TeX4ht as a label paragraph,
            # an empty paragraph, and then the actual question text.  The PDF
            # presents these as one line, so move the editable label into the
            # following text paragraph when that structure is present.
            all_paragraphs = list(document.paragraphs)  # type: ignore[attr-defined]
            paragraph_index = next(
                (index for index, item in enumerate(all_paragraphs) if item._p is paragraph._p),
                None,
            )
            if paragraph_index is None:
                continue
            next_index = paragraph_index + 1
            empty_paragraphs: list[object] = []
            while (
                next_index < len(all_paragraphs)
                and not all_paragraphs[next_index].text.strip()
                and not self._paragraph_has_embedded_object(all_paragraphs[next_index])
            ):
                empty_paragraphs.append(all_paragraphs[next_index])
                next_index += 1
            if next_index >= len(all_paragraphs):
                continue
            body = all_paragraphs[next_index]
            if self._paragraph_has_embedded_object(body):
                continue
            if re.match(r"^(?:\d+\.|\([a-z]\))", body.text.strip()):
                continue
            for empty in empty_paragraphs:
                self._delete_paragraph(empty)
            first_text_run = next((run for run in body.runs if run.text), None)
            if first_text_run is not None:
                first_text_run.text = rendered_label + " " + first_text_run.text
            else:
                prefix_run = body.add_run(rendered_label + " ")
                prefix_element = prefix_run._r
                prefix_element.getparent().remove(prefix_element)
                insert_at = 1 if body._p.pPr is not None else 0
                body._p.insert(insert_at, prefix_element)
            body.style = document.styles["dt"]  # type: ignore[attr-defined]
            for run in body.runs:
                if run.bold is None:
                    run.bold = False
            self._delete_paragraph(paragraph)

    @staticmethod
    def _restore_leading_equals(document: object) -> None:
        """Remove temporary equality markers while preserving the equals sign."""

        math_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
        text_tag = math_namespace + "t"
        math_tag = math_namespace + "oMath"
        run_tag = math_namespace + "r"
        run_properties_tag = math_namespace + "rPr"
        for paragraph in document.paragraphs:  # type: ignore[attr-defined]
            for formula in paragraph._p.iter(math_tag):
                changed = False
                for node in formula.iter(text_tag):
                    if node.text and _EQUALITY_MARKER_RE.search(node.text):
                        node.text = _EQUALITY_MARKER_RE.sub("", node.text)
                        changed = True
                if not changed:
                    continue
                # LibreOffice preserves the temporary marker as an empty
                # math run after its text is removed.  Delete that run too;
                # otherwise Word shows an editable dotted placeholder.
                for run in list(formula.iter(run_tag)):
                    content = [child for child in run if child.tag != run_properties_tag]
                    if not any(child.tag == text_tag and (child.text or "") for child in content):
                        parent = run.getparent()
                        if parent is not None:
                            parent.remove(run)

    @staticmethod
    def _restore_math_roman_styles(document: object) -> None:
        """Restore explicit ``\\mathrm``, ``\\exp``, and ``\\ln`` styling.

        TeX4ht's editable equation path emits the right OMML structure but
        does not carry ``\\mathrm`` through to the final Word run properties.
        The export-only source contains temporary markers around those spans;
        this method removes the markers and applies explicit roman/plain math
        styling to the runs between them.  Ordinary math runs are untouched.
        """

        math_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
        paragraph_tag = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
        math_tag = math_namespace + "oMath"
        run_tag = math_namespace + "r"
        run_properties_tag = math_namespace + "rPr"
        text_tag = math_namespace + "t"
        script_tag = math_namespace + "scr"
        style_tag = math_namespace + "sty"

        def set_roman_style(run: object) -> None:
            properties = run.find(run_properties_tag)
            if properties is None:
                from docx.oxml import OxmlElement  # type: ignore[import-not-found]

                properties = OxmlElement("m:rPr")
                run.insert(0, properties)
            script = properties.find(script_tag)
            if script is None:
                from docx.oxml import OxmlElement  # type: ignore[import-not-found]

                script = OxmlElement("m:scr")
                properties.append(script)
            script.set("{http://schemas.openxmlformats.org/officeDocument/2006/math}val", "roman")
            style = properties.find(style_tag)
            if style is None:
                from docx.oxml import OxmlElement  # type: ignore[import-not-found]

                style = OxmlElement("m:sty")
                properties.append(style)
            style.set("{http://schemas.openxmlformats.org/officeDocument/2006/math}val", "p")

        body = document.element.body  # type: ignore[attr-defined]
        for paragraph in body.iter(paragraph_tag):
            for formula in paragraph.iter(math_tag):
                depth = 0
                changed = False
                for run in list(formula.iter(run_tag)):
                    text_nodes = list(run.iter(text_tag))
                    run_text = "".join(node.text or "" for node in text_nodes)
                    starts = len(re.findall(r"PDFRomanS\d{3}", run_text))
                    ends = len(re.findall(r"PDFRomanE\d{3}", run_text))
                    if starts or ends:
                        changed = True
                    if depth or starts:
                        set_roman_style(run)
                    for node in text_nodes:
                        if node.text and _ROMAN_MARKER_RE.search(node.text):
                            node.text = _ROMAN_MARKER_RE.sub("", node.text)
                    depth += starts - ends
                    depth = max(0, depth)

                if not changed:
                    continue
                for run in list(formula.iter(run_tag)):
                    content = [child for child in run if child.tag != run_properties_tag]
                    if not any(child.tag == text_tag and (child.text or "") for child in content):
                        parent = run.getparent()
                        if parent is not None:
                            parent.remove(run)

    def _restore_boxed_equations(self, document: object) -> None:
        """Convert temporary markers into native OMML border boxes."""

        from docx.oxml import OxmlElement  # type: ignore[import-not-found]

        math_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
        text_tag = math_namespace + "t"
        math_tag = math_namespace + "oMath"
        for paragraph in document.paragraphs:  # type: ignore[attr-defined]
            for formula in list(paragraph._p.iter(math_tag)):
                marker_nodes = [
                    node for node in formula.iter(text_tag) if node.text and _BOX_MARKER_RE.search(node.text)
                ]
                if len(marker_nodes) < 2:
                    continue

                def direct_child(node: object) -> object | None:
                    current = node
                    while current is not None and current.getparent() is not formula:
                        current = current.getparent()
                    return current

                start_node = next((node for node in marker_nodes if "PDFBoxS" in (node.text or "")), None)
                end_node = next((node for node in marker_nodes if "PDFBoxE" in (node.text or "")), None)
                start_child = direct_child(start_node) if start_node is not None else None
                end_child = direct_child(end_node) if end_node is not None else None
                if start_child is None or end_child is None or start_child is end_child:
                    continue
                children = list(formula)
                try:
                    start_index = children.index(start_child)
                    end_index = children.index(end_child)
                except ValueError:
                    continue
                if start_index >= end_index:
                    continue

                for node in marker_nodes:
                    node.text = _BOX_MARKER_RE.sub("", node.text or "")
                border_box = OxmlElement("m:borderBox")
                expression = OxmlElement("m:e")
                for child in children[start_index + 1 : end_index]:
                    formula.remove(child)
                    expression.append(child)
                border_box.append(expression)
                for child in (start_child, end_child):
                    if child.getparent() is formula:
                        formula.remove(child)
                formula.insert(start_index, border_box)

    @staticmethod
    def _alpha_label(number: int) -> str:
        value = ""
        while number:
            number, remainder = divmod(number - 1, 26)
            value = chr(ord("a") + remainder) + value
        return f"({value})"

    @staticmethod
    def _roman_label(number: int) -> str:
        values = (
            (1000, "M"),
            (900, "CM"),
            (500, "D"),
            (400, "CD"),
            (100, "C"),
            (90, "XC"),
            (50, "L"),
            (40, "XL"),
            (10, "X"),
            (9, "IX"),
            (5, "V"),
            (4, "IV"),
            (1, "I"),
        )
        value = ""
        for amount, symbol in values:
            count, number = divmod(number, amount)
            value += symbol * count
        return f"({value})"

    def _source_list_labels(self) -> list[str]:
        """Return editable labels for TeX list items in source order."""

        token_re = re.compile(
            r"\\begin\s*\{(enumerate|parts|solution)\}\s*(?:\[([^]]*)\])?"
            r"|\\end\s*\{(enumerate|parts|solution)\}"
            r"|\\(item|part)\b",
            re.DOTALL,
        )
        stack: list[tuple[str, str]] = []
        solution_depth = 0
        solution_nested_in_parts: list[bool] = []
        labels: list[str] = []
        counters: list[int] = []
        for match in token_re.finditer(self.source_text):
            begin_environment, options, end_environment, item_kind = match.groups()
            if begin_environment:
                if begin_environment == "solution":
                    solution_depth += 1
                    solution_nested_in_parts.append(any(name == "parts" for name, _ in stack))
                    continue
                if begin_environment == "parts":
                    stack.append((begin_environment, "alpha"))
                else:
                    option_text = options or ""
                    if "Roman" in option_text:
                        style = "roman"
                    elif "alph" in option_text or "alpha" in option_text:
                        style = "alpha"
                    else:
                        style = "arabic"
                    stack.append((begin_environment, style))
                counters.append(0)
                continue
            if end_environment:
                if end_environment == "solution":
                    solution_depth = max(0, solution_depth - 1)
                    if solution_nested_in_parts:
                        solution_nested_in_parts.pop()
                elif stack:
                    stack.pop()
                    counters.pop()
                continue
            if not stack or not counters:
                continue
            environment, style = stack[-1]
            if item_kind == "part":
                if (
                    environment != "parts"
                    or solution_depth == 0
                    or not solution_nested_in_parts[-1]
                ):
                    continue
            elif item_kind == "item":
                if environment != "enumerate":
                    continue
            else:
                continue
            counters[-1] += 1
            number = counters[-1]
            if style == "roman":
                labels.append(self._roman_label(number))
            elif style == "alpha":
                labels.append(self._alpha_label(number))
            else:
                labels.append(f"{number}.")
        return labels

    def _restore_list_labels(self, document: object) -> None:
        """Replace TeX4ht's presentation-only list labels with Word text."""

        from docx.oxml.ns import qn  # type: ignore[import-not-found]
        from docx.shared import Inches  # type: ignore[import-not-found]

        source_labels = self._source_list_labels()
        label_index = 0
        number = 0
        in_list = False
        for paragraph in document.paragraphs:  # type: ignore[attr-defined]
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name != "Inside-enumerate":
                in_list = False
                continue
            number = number + 1 if in_list else 1
            in_list = True
            # Remove the numbering style and preserve the visual indent as
            # direct formatting so the number can be edited like normal text.
            paragraph.style = document.styles["dd"]  # type: ignore[attr-defined]
            properties = paragraph._p.pPr
            if properties is not None:
                numbering = properties.find(qn("w:numPr"))
                if numbering is not None:
                    properties.remove(numbering)
            paragraph.paragraph_format.left_indent = Inches(0.45)
            paragraph.paragraph_format.first_line_indent = Inches(-0.2)
            if label_index < len(source_labels):
                label = source_labels[label_index]
            else:
                label = f"{number}."
            label_index += 1
            prefix_run = paragraph.add_run(f"{label} ")
            prefix_element = prefix_run._r
            prefix_element.getparent().remove(prefix_element)
            insert_at = 1 if paragraph._p.pPr is not None else 0
            paragraph._p.insert(insert_at, prefix_element)

    @staticmethod
    def _frame_table(table: object) -> None:
        from docx.oxml import OxmlElement  # type: ignore[import-not-found]
        from docx.oxml.ns import qn  # type: ignore[import-not-found]

        table.autofit = True  # type: ignore[attr-defined]
        properties = table._tbl.tblPr  # type: ignore[attr-defined]
        borders = properties.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            properties.append(borders)
        for side in ("top", "left", "bottom", "right"):
            element = borders.find(qn(f"w:{side}"))
            if element is None:
                element = OxmlElement(f"w:{side}")
                borders.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "6")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "000000")
        for cell in table.rows[0].cells:  # type: ignore[attr-defined]
            cell_properties = cell._tc.get_or_add_tcPr()
            margins = cell_properties.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell_properties.append(margins)
            for side, value in (("top", 80), ("bottom", 80), ("left", 120), ("right", 120)):
                margin = margins.find(qn(f"w:{side}"))
                if margin is None:
                    margin = OxmlElement(f"w:{side}")
                    margins.append(margin)
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")

    def _restore_solution_frames(self, document: object) -> None:
        """Keep TeX4ht's solution paragraphs inside their editable frame."""

        from docx.table import Table  # type: ignore[import-not-found]
        from docx.text.paragraph import Paragraph  # type: ignore[import-not-found]

        body = document._body._body  # type: ignore[attr-defined]
        body_children = list(body.iterchildren())
        solution_tables: list[object] = []
        for child in body_children:
            if not child.tag.endswith("}tbl"):
                continue
            table = Table(child, document)
            cells = [cell for row in table.rows for cell in row.cells]
            if any(
                any(paragraph.text.strip().startswith("Solution:") for paragraph in cell.paragraphs)
                for cell in cells
            ):
                solution_tables.append(table)

        for table in solution_tables:
            cell = next(
                (
                    cell
                    for row in table.rows
                    for cell in row.cells
                    if any(paragraph.text.strip().startswith("Solution:") for paragraph in cell.paragraphs)
                ),
                None,
            )
            if cell is None:
                continue
            body_children = list(body.iterchildren())
            current_index = body_children.index(table._tbl)  # type: ignore[attr-defined]
            moved: list[object] = []
            for child in body_children[current_index + 1 :]:
                if child.tag.endswith("}p"):
                    paragraph = Paragraph(child, document)
                    if re.match(r"^\d+\.\s*\[\d+ points\](?:\s|$)", paragraph.text.strip()):
                        break
                    moved.append(child)
                elif child.tag.endswith("}tbl"):
                    break
            for child in moved:
                cell._tc.append(child)

        # If LibreOffice received the repaired ODT without solution tables,
        # create the frames from the source-level solution boundaries.  The
        # question label restored above is the reliable end marker; numeric
        # list labels are not boundaries because solution parts can use them.
        paragraphs = list(document.paragraphs)  # type: ignore[attr-defined]
        starts = [
            index
            for index, paragraph in enumerate(paragraphs)
            if paragraph.text.strip().startswith("Solution:")
        ]
        for start in reversed(starts):
            end = len(paragraphs)
            for index in range(start + 1, len(paragraphs)):
                if re.match(r"^\d+\.\s*\[\d+ points\](?:\s|$)", paragraphs[index].text.strip()):
                    end = index
                    break
            if end <= start:
                continue
            table = document.add_table(rows=1, cols=1)  # type: ignore[attr-defined]
            self._frame_table(table)
            table_element = table._tbl
            body = document._body._body  # type: ignore[attr-defined]
            body.remove(table_element)
            body.insert(body.index(paragraphs[start]._p), table_element)
            cell = table.cell(0, 0)
            for index in range(start, end):
                cell._tc.append(paragraphs[index]._p)


    @staticmethod
    def _add_page_numbers(document: object) -> None:
        """Add editable ``Page N`` fields to every Word section footer."""

        from docx.oxml import OxmlElement  # type: ignore[import-not-found]
        from docx.oxml.ns import qn  # type: ignore[import-not-found]
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-not-found]

        seen_footers: set[int] = set()
        for section in document.sections:  # type: ignore[attr-defined]
            footer = section.footer
            identity = id(footer._element)
            if identity in seen_footers:
                continue
            seen_footers.add(identity)
            paragraph = footer.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in list(paragraph.runs):
                run._element.getparent().remove(run._element)
            paragraph.add_run("Page ")
            run = paragraph.add_run()
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = " PAGE "
            separate = OxmlElement("w:fldChar")
            separate.set(qn("w:fldCharType"), "separate")
            displayed = OxmlElement("w:t")
            displayed.text = "1"
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            run._r.extend((begin, instruction, separate, displayed, end))

        settings = document.settings.element  # type: ignore[attr-defined]
        update_fields = settings.find(qn("w:updateFields"))
        if update_fields is None:
            update_fields = OxmlElement("w:updateFields")
            settings.append(update_fields)
        update_fields.set(qn("w:val"), "true")

    def _insert_figures(self, document: object, project_dir: Path) -> None:
        """Restore source figures that LibreOffice omits from TeX4ht ODT."""

        figure_names = _graphics_names(self.source_text)
        if not figure_names:
            return
        figure_paths = [_find_graphic(project_dir, name) for name in figure_names]
        if not all(figure_paths):
            LOGGER.warning("Some LaTeX figures were not found for DOCX insertion: %s", figure_names)
            return

        try:
            from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-not-found]
            from docx.oxml import OxmlElement  # type: ignore[import-not-found]
            from docx.shared import Inches, Pt  # type: ignore[import-not-found]
            from docx.text.paragraph import Paragraph  # type: ignore[import-not-found]
        except ImportError as exc:
            raise LatexWorkspaceError(
                "python-docx is required to place source figures in the editable DOCX.",
                "Install the project requirements, then try editable Word export again.",
            ) from exc

        anchors = [
            paragraph
            for paragraph in document.paragraphs
            if "shown below" in paragraph.text
            or ("shown" in paragraph.text and "Figure" in paragraph.text)
        ]
        if len(anchors) < len(figure_paths):
            # Some TeX4ht versions split or simplify the figure sentence.
            # Use a nearby figure reference as a final source-aware anchor.
            anchors = [
                paragraph
                for paragraph in document.paragraphs
                if "Figure" in paragraph.text or "figure" in paragraph.text
            ]
        if len(anchors) < len(figure_paths):
            LOGGER.warning("Could not locate all figure anchors in the generated DOCX.")
            return

        widths = _graphics_widths(self.source_text)
        for anchor, figure_path in zip(anchors, figure_paths):
            paragraph_element = OxmlElement("w:p")
            anchor._p.addnext(paragraph_element)
            paragraph = Paragraph(paragraph_element, anchor._parent)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(6)
            width = widths.get(figure_path.stem, 3.575)
            paragraph.add_run().add_picture(str(figure_path), width=Inches(width))
    def close(self) -> None:
        return None


def _compile_command(main_name: str, engine: str, build_dir: Path) -> list[str]:
    if engine == "auto" and _command_path("latexmk"):
        return [
            "latexmk",
            "-pdf",
            "-interaction=nonstopmode",
            "-halt-on-error",
            "-file-line-error",
            "-outdir=" + str(build_dir),
            main_name,
        ]

    selected = engine
    if selected == "auto":
        selected = "pdflatex"
    if not _command_path(selected):
        raise LatexWorkspaceError(
            f"The selected TeX compiler was not found: {selected}",
            f"Install {selected}, or choose a compiler that is already installed.",
        )
    return [
        selected,
        "-interaction=nonstopmode",
        "-halt-on-error",
        "-file-line-error",
        "-output-directory=" + str(build_dir),
        main_name,
    ]


def compile_latex(
    tex_path: Path,
    source_text: str,
    reference_pdf: Path | None = None,
    engine: str = "auto",
) -> CompileResult:
    """Compile a source buffer in an isolated temporary project directory."""

    if not tex_path.is_file():
        raise LatexWorkspaceError(
            f"The TeX source does not exist: {tex_path}",
            "Choose an existing .tex source file first.",
        )

    temp_root = Path(tempfile.mkdtemp(prefix="pdf_to_word_latex_"))
    project_dir = temp_root / "project"
    build_dir = temp_root / "build"
    project_dir.mkdir()
    build_dir.mkdir()

    standalone = _is_standalone_tex(source_text)
    setup_warnings: list[str] = []
    if standalone:
        main_name = tex_path.name
        main_path = project_dir / main_name
        main_path.write_text(source_text, encoding="utf-8")
    else:
        body_name = "source_body.tex"
        (project_dir / body_name).write_text(source_text, encoding="utf-8")
        main_name = "main.tex"
        template_path = _find_master_template(tex_path, source_text)
        if template_path is None:
            wrapper = _wrapper_source(body_name)
        else:
            wrapper, setup_warnings = _template_wrapper_source(
                template_path,
                body_name,
                project_dir,
            )
        (project_dir / main_name).write_text(wrapper, encoding="utf-8")

    _copy_existing_graphics(tex_path.parent, project_dir, source_text)
    recovered_assets = _recover_graphics_from_pdf(source_text, reference_pdf, project_dir)

    command = _compile_command(main_name, engine, build_dir)
    log_parts: list[str] = setup_warnings.copy()
    try:
        first = subprocess.run(
            command,
            cwd=str(project_dir),
            capture_output=True,
            text=True,
            check=False,
        )
        log_parts.extend([first.stdout, first.stderr])

        # A direct engine needs a second pass for references and page layout.
        if first.returncode == 0 and command[0] != "latexmk":
            second = subprocess.run(
                command,
                cwd=str(project_dir),
                capture_output=True,
                text=True,
                check=False,
            )
            log_parts.extend([second.stdout, second.stderr])
            first = second
    except OSError as exc:
        raise LatexWorkspaceError(
            f"Could not start the TeX compiler: {exc}",
            "The TeX compiler could not be started. Check the compiler selection and your TeX installation.",
        ) from exc

    pdf_path = build_dir / Path(main_name).with_suffix(".pdf").name
    log = "\n".join(part for part in log_parts if part)
    if first.returncode == 0 and pdf_path.is_file():
        log = _filter_resolved_reference_warnings(log, build_dir / Path(main_name).with_suffix(".aux").name)
    if first.returncode != 0 or not pdf_path.is_file():
        raise LatexWorkspaceError(
            f"LaTeX compilation failed for {tex_path}:\n{log}",
            "LaTeX could not compile this source. Review the compiler log below for the line that needs attention.",
        )
    return CompileResult(pdf_path, log, recovered_assets, build_dir)


class LatexWorkspaceApp:
    """Tkinter UI for source editing, live PDF preview, and Word export."""

    def __init__(self, root: tk.Tk, initial_tex: str | None = None) -> None:
        self.root = root
        self.root.title("LaTeX Workspace")
        self.root.geometry("1440x900")
        self.root.minsize(980, 650)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

        self.tex_path: Path | None = None
        self.reference_pdf: Path | None = None
        self.preview_pdf: Path | None = None
        self.preview_pages = 0
        self.preview_page_index = 0
        self.preview_image: tk.PhotoImage | None = None
        self.compile_result: CompileResult | None = None
        self.last_compiled_text = ""
        self.compile_thread: threading.Thread | None = None
        self.job_thread: threading.Thread | None = None
        self.compile_queue: queue.Queue[tuple[str, object]] = queue.Queue()
        self.auto_compile = tk.BooleanVar(value=True)
        self.engine = tk.StringVar(value="auto")
        self.status_var = tk.StringVar(value="Open a .tex file to begin")
        self.path_var = tk.StringVar(value="No source selected")
        self.reference_var = tk.StringVar(value="Reference PDF: automatic")
        self.page_var = tk.StringVar(value="No preview")
        self._compile_timer: str | None = None

        self._build_ui()
        self.root.after(100, self._poll_queue)
        if initial_tex:
            self.root.after(100, lambda: self.open_tex(Path(initial_tex)))

    def _build_ui(self) -> None:
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(1, weight=1)

        toolbar = ttk.Frame(self.root, padding=(10, 8))
        toolbar.grid(row=0, column=0, sticky="ew")
        for column in range(8):
            toolbar.columnconfigure(column, weight=1 if column == 7 else 0)

        ttk.Button(toolbar, text="Open .tex", command=self._choose_tex).grid(row=0, column=0, padx=(0, 6))
        ttk.Button(toolbar, text="Save Source", command=self.save_source).grid(row=0, column=1, padx=6)
        ttk.Button(toolbar, text="Compile / Preview", command=self.start_compile).grid(row=0, column=2, padx=6)
        ttk.Checkbutton(toolbar, text="Auto compile", variable=self.auto_compile).grid(row=0, column=3, padx=6)
        ttk.Label(toolbar, text="Engine:").grid(row=0, column=4, padx=(12, 3))
        ttk.Combobox(
            toolbar,
            textvariable=self.engine,
            values=("auto", "pdflatex", "xelatex", "lualatex"),
            state="readonly",
            width=10,
        ).grid(row=0, column=5, padx=3)
        ttk.Button(toolbar, text="Set Reference PDF", command=self._choose_reference_pdf).grid(row=0, column=6, padx=6)
        ttk.Label(toolbar, textvariable=self.status_var).grid(row=0, column=7, sticky="e", padx=(12, 0))

        content = ttk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        content.grid(row=1, column=0, sticky="nsew", padx=10)

        source_frame = ttk.Frame(content, padding=(0, 0, 6, 0))
        source_frame.rowconfigure(2, weight=1)
        source_frame.columnconfigure(0, weight=1)
        ttk.Label(source_frame, textvariable=self.path_var).grid(row=0, column=0, sticky="w")
        ttk.Label(source_frame, text="Edit the original LaTeX source. The preview compiles a temporary copy.").grid(
            row=1, column=0, sticky="w", pady=(4, 6)
        )
        source_text_frame = ttk.Frame(source_frame)
        source_text_frame.grid(row=2, column=0, sticky="nsew")
        source_text_frame.rowconfigure(0, weight=1)
        source_text_frame.columnconfigure(0, weight=1)
        self.editor = tk.Text(
            source_text_frame,
            undo=True,
            wrap="none",
            font=("Menlo", 12),
            padx=10,
            pady=10,
        )
        self.editor.grid(row=0, column=0, sticky="nsew")
        y_scroll = ttk.Scrollbar(source_text_frame, orient="vertical", command=self.editor.yview)
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll = ttk.Scrollbar(source_text_frame, orient="horizontal", command=self.editor.xview)
        x_scroll.grid(row=1, column=0, sticky="ew")
        self.editor.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        self.editor.bind("<KeyRelease>", self._on_source_changed)
        content.add(source_frame, weight=1)

        preview_frame = ttk.Frame(content, padding=(6, 0, 0, 0))
        preview_frame.rowconfigure(1, weight=1)
        preview_frame.columnconfigure(0, weight=1)
        preview_toolbar = ttk.Frame(preview_frame)
        preview_toolbar.grid(row=0, column=0, sticky="ew", pady=(0, 6))
        ttk.Button(preview_toolbar, text="Previous", command=self._previous_page).pack(side="left")
        ttk.Button(preview_toolbar, text="Next", command=self._next_page).pack(side="left", padx=6)
        ttk.Label(preview_toolbar, textvariable=self.page_var).pack(side="left", padx=8)
        ttk.Label(preview_toolbar, textvariable=self.reference_var).pack(side="right")

        self.preview_canvas = tk.Canvas(preview_frame, background="#d7d7d7", highlightthickness=0)
        self.preview_canvas.grid(row=1, column=0, sticky="nsew")
        preview_scroll = ttk.Scrollbar(preview_frame, orient="vertical", command=self.preview_canvas.yview)
        preview_scroll.grid(row=1, column=1, sticky="ns")
        self.preview_canvas.configure(yscrollcommand=preview_scroll.set)
        content.add(preview_frame, weight=1)

        bottom = ttk.Frame(self.root, padding=(10, 8))
        bottom.grid(row=2, column=0, sticky="ew")
        bottom.columnconfigure(0, weight=1)
        ttk.Label(bottom, text="Compiler log").grid(row=0, column=0, sticky="w")
        self.log_text = tk.Text(bottom, height=6, wrap="word", state="disabled", font=("Menlo", 10))
        self.log_text.grid(row=1, column=0, sticky="ew", pady=(4, 8))
        action_bar = ttk.Frame(bottom)
        action_bar.grid(row=2, column=0, sticky="ew")
        ttk.Button(action_bar, text="Export Exact Word", command=lambda: self.export_word(exact=True)).pack(side="left")
        ttk.Button(action_bar, text="Export Editable Word", command=lambda: self.export_word(exact=False)).pack(
            side="left", padx=8
        )
        ttk.Button(action_bar, text="Open Preview PDF", command=self.open_preview_pdf).pack(side="left")

    def _choose_tex(self) -> None:
        path = filedialog.askopenfilename(
            parent=self.root,
            title="Open LaTeX source",
            filetypes=[("LaTeX source", "*.tex"), ("All files", "*.*")],
        )
        if path:
            self.open_tex(Path(path))

    def open_tex(self, path: Path) -> None:
        try:
            source = path.read_text(encoding="utf-8")
        except (OSError, UnicodeError) as exc:
            messagebox.showerror("Could not open source", str(exc), parent=self.root)
            return
        self.tex_path = path.resolve()
        self.editor.delete("1.0", "end")
        self.editor.insert("1.0", source)
        self.last_compiled_text = ""
        self.reference_pdf = _find_reference_pdf(self.tex_path)
        self.path_var.set(str(self.tex_path))
        self.reference_var.set(
            f"Reference PDF: {self.reference_pdf.name}" if self.reference_pdf else "Reference PDF: automatic"
        )
        self._append_log(f"Opened {self.tex_path}\n")
        self.status_var.set("Ready to compile")
        self._schedule_compile()

    def _choose_reference_pdf(self) -> None:
        path = filedialog.askopenfilename(
            parent=self.root,
            title="Choose a reference PDF for missing figures",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")],
        )
        if path:
            self.reference_pdf = Path(path).resolve()
            self.reference_var.set(f"Reference PDF: {self.reference_pdf.name}")
            self._schedule_compile()

    def _on_source_changed(self, _event: tk.Event) -> None:
        if self.auto_compile.get():
            self._schedule_compile()

    def _schedule_compile(self) -> None:
        if self._compile_timer is not None:
            self.root.after_cancel(self._compile_timer)
        self._compile_timer = self.root.after(1200, self.start_compile)

    def save_source(self) -> None:
        if self.tex_path is None:
            messagebox.showinfo("No source", "Open a .tex file first.", parent=self.root)
            return
        try:
            self.tex_path.write_text(self.editor.get("1.0", "end-1c"), encoding="utf-8")
        except OSError as exc:
            messagebox.showerror("Could not save source", str(exc), parent=self.root)
            return
        self.status_var.set("Source saved")
        self._append_log(f"Saved {self.tex_path}\n")

    def start_compile(self) -> None:
        self._compile_timer = None
        if self.tex_path is None:
            return
        if self.compile_thread and self.compile_thread.is_alive():
            return
        source_text = self.editor.get("1.0", "end-1c")
        self.status_var.set("Compiling LaTeX...")
        self.compile_thread = threading.Thread(
            target=self._compile_worker,
            args=(self.tex_path, source_text, self.reference_pdf, self.engine.get()),
            daemon=True,
        )
        self.compile_thread.start()

    def _compile_worker(self, tex_path: Path, source_text: str, reference_pdf: Path | None, engine: str) -> None:
        try:
            result = compile_latex(tex_path, source_text, reference_pdf, engine)
            self.compile_queue.put(("compile_ok", result))
        except Exception as exc:
            self.compile_queue.put(("compile_error", exc))

    def _poll_queue(self) -> None:
        try:
            while True:
                kind, payload = self.compile_queue.get_nowait()
                if kind == "compile_ok":
                    self._compile_finished(payload)  # type: ignore[arg-type]
                elif kind == "compile_error":
                    self._compile_failed(payload)  # type: ignore[arg-type]
                elif kind == "status":
                    self.status_var.set(str(payload))
                elif kind == "export_ok":
                    self._export_finished(payload)  # type: ignore[arg-type]
                elif kind == "export_error":
                    self._export_failed(payload)  # type: ignore[arg-type]
        except queue.Empty:
            pass
        self.root.after(100, self._poll_queue)

    def _compile_finished(self, result: CompileResult) -> None:
        self.compile_result = result
        self.preview_pdf = result.pdf_path
        self.preview_page_index = 0
        self.last_compiled_text = self.editor.get("1.0", "end-1c")
        self.status_var.set("Compiled successfully")
        self._append_log(result.log)
        if result.recovered_assets:
            self._append_log("Recovered figures from reference PDF: " + ", ".join(result.recovered_assets) + "\n")
        self._load_preview_page()

    def _compile_failed(self, exc: Exception) -> None:
        self.status_var.set("Compilation failed")
        self._append_log(str(exc))
        if isinstance(exc, LatexWorkspaceError):
            self._append_log(f"\n{exc.user_message}\n")

    def _load_preview_page(self) -> None:
        if self.preview_pdf is None:
            return
        try:
            import pymupdf as fitz  # type: ignore[import-not-found]
        except ImportError:
            try:
                import fitz  # type: ignore[import-not-found]
            except ImportError:
                self._append_log("PyMuPDF is required for the PDF preview.\n")
                return
        try:
            document = fitz.open(str(self.preview_pdf))
            with document:
                self.preview_pages = document.page_count
                if self.preview_pages == 0:
                    return
                page = document.load_page(self.preview_page_index)
                available_width = max(self.preview_canvas.winfo_width() - 28, 500)
                scale = min(1.35, available_width / float(page.rect.width))
                pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                encoded = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
            self.preview_image = tk.PhotoImage(data=encoded)
            self.preview_canvas.delete("all")
            self.preview_canvas.create_image(14, 14, anchor="nw", image=self.preview_image)
            self.preview_canvas.configure(
                scrollregion=(0, 0, self.preview_image.width() + 28, self.preview_image.height() + 28)
            )
            self.page_var.set(f"Page {self.preview_page_index + 1} of {self.preview_pages}")
        except Exception as exc:
            self._append_log(f"Could not render preview: {exc}\n")

    def _previous_page(self) -> None:
        if self.preview_pages and self.preview_page_index > 0:
            self.preview_page_index -= 1
            self._load_preview_page()

    def _next_page(self) -> None:
        if self.preview_pages and self.preview_page_index + 1 < self.preview_pages:
            self.preview_page_index += 1
            self._load_preview_page()

    def open_preview_pdf(self) -> None:
        if self.preview_pdf and self.preview_pdf.is_file():
            self._open_path(self.preview_pdf)

    def export_word(self, exact: bool) -> None:
        if self.preview_pdf is None or self.compile_result is None:
            messagebox.showinfo("Compile first", "Compile the LaTeX source before exporting Word.", parent=self.root)
            return
        current_text = self.editor.get("1.0", "end-1c")
        if current_text != self.last_compiled_text:
            messagebox.showinfo("Compile first", "The source has changed since the last preview. Compile it again before exporting.", parent=self.root)
            return
        default_name = (self.tex_path.stem if self.tex_path else "document") + ".docx"
        output = filedialog.asksaveasfilename(
            parent=self.root,
            title="Export Word document",
            initialfile=default_name,
            initialdir=str(self.tex_path.parent if self.tex_path else Path.cwd()),
            defaultextension=".docx",
            filetypes=[("Word document", "*.docx")],
        )
        if not output:
            return
        output_path = Path(output)
        if output_path.exists() and not messagebox.askyesno("Overwrite file", f"Replace {output_path.name}?", parent=self.root):
            return

        self.status_var.set("Exporting Word document...")
        self.job_thread = threading.Thread(
            target=self._export_worker,
            args=(self.preview_pdf, output_path, exact, self.tex_path, current_text, self.compile_result),
            daemon=True,
        )
        self.job_thread.start()

    def _export_worker(
        self,
        pdf_path: Path,
        output_path: Path,
        exact: bool,
        tex_path: Path | None,
        source_text: str,
        compile_result: CompileResult,
    ) -> None:
        try:
            if exact:
                converter = VisualFidelityConverter()
                mode = "exact appearance"
            else:
                if tex_path is None:
                    raise LatexWorkspaceError(
                        "No LaTeX source is attached to this preview.",
                        "Open and compile a .tex source before exporting editable Word.",
                    )
                converter = LatexSourceDocxConverter(compile_result, tex_path, source_text)
                mode = "editable LaTeX-source Word"
            try:
                result = converter.convert_pdf(
                    pdf_path,
                    output_path,
                    overwrite=True,
                    status_callback=self._thread_status,
                )
            finally:
                converter.close()
            self.compile_queue.put(("export_ok", ExportResult(result, mode)))
        except Exception as exc:
            self.compile_queue.put(("export_error", exc))

    def _thread_status(self, message: str) -> None:
        self.compile_queue.put(("status", message))

    def _export_finished(self, result: ExportResult) -> None:
        self.status_var.set("Word export complete")
        self._append_log(f"Created {result.mode}: {result.output_path}\n")
        if messagebox.askyesno("Export complete", f"Open {Path(result.output_path).name}?", parent=self.root):
            self._open_path(Path(result.output_path))

    def _export_failed(self, exc: Exception) -> None:
        self.status_var.set("Word export failed")
        message = getattr(exc, "user_message", str(exc))
        self._append_log(f"Export failed: {exc}\n")
        messagebox.showerror("Word export failed", message, parent=self.root)

    def _append_log(self, text: str) -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", text.rstrip() + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    @staticmethod
    def _open_path(path: Path) -> None:
        try:
            if platform.system() == "Darwin":
                subprocess.Popen(["open", str(path)])
            elif platform.system() == "Windows":
                os.startfile(str(path))  # type: ignore[attr-defined]
            else:
                subprocess.Popen(["xdg-open", str(path)])
        except OSError as exc:
            LOGGER.warning("Could not open %s: %s", path, exc)

    def _on_close(self) -> None:
        self.root.destroy()


def main() -> None:
    root = tk.Tk()
    initial_tex = sys.argv[1] if len(sys.argv) > 1 else None
    LatexWorkspaceApp(root, initial_tex=initial_tex)
    root.mainloop()


if __name__ == "__main__":
    main()
