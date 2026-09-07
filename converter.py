"""High-fidelity PDF-to-DOCX conversion and command-line entry point."""

from __future__ import annotations

import argparse
import logging
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
from typing import Callable

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


LOGGER = logging.getLogger(__name__)
POINTS_PER_INCH = 72.0
StatusCallback = Callable[[str], None]


class PDFConversionError(Exception):
    """Base exception for expected conversion failures."""


class PDFValidationError(PDFConversionError):
    """The input or output path is invalid."""


class OutputExistsError(PDFConversionError):
    """The requested output exists and overwrite was not approved."""


class ConversionCancelledError(PDFConversionError):
    """Conversion was cancelled before completion."""


def _absolute_path(value: str | os.PathLike[str], description: str) -> Path:
    try:
        return Path(os.path.abspath(os.fspath(value)))
    except (TypeError, ValueError, OSError) as exc:
        raise PDFValidationError(f"Invalid {description} path: {value!r}") from exc


def normalise_pdf_path(pdf_path: str | os.PathLike[str]) -> Path:
    """Validate and normalize a PDF input path."""

    path = _absolute_path(pdf_path, "PDF")
    if not path.exists() or not path.is_file():
        raise PDFValidationError(f"PDF does not exist or is not a file: {path}")
    if path.suffix.lower() != ".pdf":
        raise PDFValidationError(f"Input is not a PDF: {path}")
    return path


def normalise_output_path(
    pdf_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str] | None = None,
) -> Path:
    """Return an absolute DOCX path, deriving it from the PDF when omitted."""

    input_path = _absolute_path(pdf_path, "PDF")
    if output_path is None or not str(output_path).strip():
        return input_path.with_suffix(".docx")

    path = _absolute_path(output_path, "output")
    if path.suffix == "":
        path = path.with_suffix(".docx")
    if path.suffix.lower() != ".docx":
        raise PDFValidationError(f"Output is not a DOCX path: {path}")
    return path


def _configure_section(section, width_in: float, height_in: float) -> None:
    section.page_width = Inches(width_in)
    section.page_height = Inches(height_in)
    section.top_margin = Inches(0)
    section.bottom_margin = Inches(0)
    section.left_margin = Inches(0)
    section.right_margin = Inches(0)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)
    section.gutter = Inches(0)


def _configure_page_paragraph(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    formatting = paragraph.paragraph_format
    formatting.space_before = Pt(0)
    formatting.space_after = Pt(0)
    formatting.left_indent = Inches(0)
    formatting.right_indent = Inches(0)
    formatting.first_line_indent = Inches(0)
    formatting.line_spacing = 1


def convert_pdf_to_docx(
    pdf_path: str | os.PathLike[str],
    output_path: str | os.PathLike[str] | None = None,
    *,
    dpi: int = 300,
    overwrite: bool = False,
    status_callback: StatusCallback | None = None,
    cancel_event: Event | None = None,
) -> Path:
    """Render a PDF into a visually faithful Word document.

    Each source page becomes one high-resolution inline image on a matching
    marginless Word page. The temporary page images are removed automatically.
    """

    pdf = normalise_pdf_path(pdf_path)
    docx = normalise_output_path(pdf, output_path)

    if dpi < 72 or dpi > 600:
        raise PDFValidationError("Rendering DPI must be between 72 and 600.")
    if docx.exists() and not overwrite:
        raise OutputExistsError(f"Output already exists: {docx}")
    if not docx.parent.exists():
        raise PDFValidationError(f"Output folder does not exist: {docx.parent}")
    if not os.access(docx.parent, os.W_OK):
        raise PDFValidationError(f"Output folder is not writable: {docx.parent}")

    if cancel_event and cancel_event.is_set():
        raise ConversionCancelledError("Conversion cancelled before opening the PDF.")

    pdf_document = pymupdf.open(pdf)
    try:
        if pdf_document.page_count == 0:
            raise PDFValidationError(f"The PDF has no pages: {pdf}")

        word_document = Document()
        scale = dpi / POINTS_PER_INCH
        total_pages = pdf_document.page_count

        with TemporaryDirectory(prefix="pdf_to_word_") as temp_dir:
            temp_root = Path(temp_dir)
            for page_number, page in enumerate(pdf_document):
                if cancel_event and cancel_event.is_set():
                    raise ConversionCancelledError("Conversion cancelled.")

                page_label = f"Rendering page {page_number + 1} of {total_pages}..."
                if status_callback:
                    status_callback(page_label)

                width_in = float(page.rect.width) / POINTS_PER_INCH
                height_in = float(page.rect.height) / POINTS_PER_INCH
                if page_number == 0:
                    section = word_document.sections[0]
                else:
                    section = word_document.add_section(WD_SECTION.NEW_PAGE)
                _configure_section(section, width_in, height_in)

                paragraph = word_document.add_paragraph()
                _configure_page_paragraph(paragraph)

                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(scale, scale),
                    alpha=False,
                    annots=True,
                )
                image_path = temp_root / f"page-{page_number + 1:04d}.png"
                pixmap.save(image_path)
                paragraph.add_run().add_picture(
                    str(image_path), width=Inches(width_in), height=Inches(height_in)
                )

        if cancel_event and cancel_event.is_set():
            raise ConversionCancelledError("Conversion cancelled before saving.")

        properties = word_document.core_properties
        properties.title = pdf.stem
        properties.subject = "High-fidelity PDF to Word conversion"
        properties.author = ""
        word_document.save(docx)
    finally:
        pdf_document.close()

    if status_callback:
        status_callback("Conversion complete.")
    return docx


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Convert a PDF into a visually faithful Word document."
    )
    parser.add_argument("input_pdf", type=Path, help="Path to the source PDF")
    parser.add_argument(
        "--output",
        "-o",
        type=Path,
        help="Output DOCX path; defaults to the input filename with a .docx suffix",
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=300,
        help="Rendering resolution in DPI; default: 300",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Replace an existing DOCX at the output path",
    )
    return parser


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    args = build_parser().parse_args()
    try:
        output = convert_pdf_to_docx(
            args.input_pdf,
            args.output,
            dpi=args.dpi,
            overwrite=args.overwrite,
            status_callback=lambda message: LOGGER.info(message),
        )
    except PDFConversionError as exc:
        raise SystemExit(f"Conversion failed: {exc}") from exc
    print(f"Created {output}")


if __name__ == "__main__":
    main()
